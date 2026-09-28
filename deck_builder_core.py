import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ==========================================
# PALETTE & DESIGN SYSTEM (Canva & Google Inspired)
# ==========================================
NAVY_DEEP     = RGBColor(10, 22, 40)      # #0A1628
NAVY_CARD     = RGBColor(18, 32, 58)      # #12203A
NAVY_LIGHT    = RGBColor(30, 48, 80)      # #1E3050
EMERALD       = RGBColor(16, 185, 129)    # #10B981 (Primary Brand)
EMERALD_DARK  = RGBColor(5, 150, 105)     # #059669
EMERALD_PALE  = RGBColor(236, 253, 245)   # #ECFDF5
TEAL          = RGBColor(13, 148, 136)    # #0D9488
SLATE_BG      = RGBColor(248, 250, 252)   # #F8FAFC (Clean background)
WHITE         = RGBColor(255, 255, 255)
CARD_BG       = RGBColor(255, 255, 255)   # White card
BORDER_LIGHT  = RGBColor(226, 232, 240)   # #E2E8F0
TEXT_DARK     = RGBColor(15, 23, 42)      # #0F172A
TEXT_MUTED    = RGBColor(100, 116, 139)   # #64748B
TEXT_LIGHT    = RGBColor(241, 245, 249)   # #F1F5F9
TEXT_SUBTLE   = RGBColor(148, 163, 184)   # #94A3B8
AMBER         = RGBColor(217, 119, 6)     # #D97706
AMBER_BG      = RGBColor(254, 243, 199)   # #FEF3C7
RED_CRIT      = RGBColor(225, 29, 72)     # #E11D48
RED_BG        = RGBColor(255, 241, 242)   # #FFF1F2

FONT_HEADING = "Helvetica"
FONT_BODY    = "Helvetica"

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def set_slide_background(slide, color):
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = color
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, tag_text, title_text, subtitle_text=None, is_dark=False):
    # Tag Pill
    tag_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(3.8), Inches(0.32)
    )
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = EMERALD if is_dark else EMERALD_PALE
    tag_box.line.color.rgb = EMERALD
    tag_box.line.width = Pt(1)
    tf_tag = tag_box.text_frame
    tf_tag.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.alignment = PP_ALIGN.CENTER
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = NAVY_DEEP if is_dark else EMERALD_DARK

    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.92), Inches(11.8), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE if is_dark else TEXT_DARK

    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.48), Inches(11.8), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = TEXT_SUBTLE if is_dark else TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_LIGHT, border_width=1):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

print("Helper core ready!")
